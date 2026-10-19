import io
import logging
import mimetypes
from dataclasses import dataclass

from app.llm import LLMError, LLMProvider

log = logging.getLogger(__name__)


class UnsupportedSource(ValueError):
    pass


@dataclass
class ExtractedSource:
    kind: str
    name: str | None
    text: str


TEXT_EXT = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".log"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".mp3", ".wav", ".m4a"}

_UNREADABLE = {
    "pdf": "Couldn't read this PDF. It may be damaged or password-protected.",
    "docx": "Couldn't read this Word document. It may be damaged, password-protected or not a real .docx file.",
}
_NO_TEXT = {
    "pdf": "This PDF has no text that can be extracted; it may be a scanned image. "
    "Paste the text, or upload the pages as images instead.",
    "docx": "This Word document has no text to work from. Paste the text instead.",
}


def extract(filename: str, data: bytes, provider: LLMProvider) -> ExtractedSource:
    """Blocking: image and video extraction call the model, so run it in a thread."""
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext in (".pdf", ".docx"):
        return _document(ext[1:], filename, data)
    if ext in TEXT_EXT:
        return ExtractedSource("text", filename, data.decode("utf-8", errors="replace"))
    if ext in IMAGE_EXT:
        mime = mimetypes.guess_type(filename)[0] or "image/png"
        try:
            text = provider.describe_image(data, mime)
        except LLMError as e:
            raise UnsupportedSource(f"Could not read the image: {e}") from e
        return ExtractedSource("image", filename, f"[Image content extracted by vision model]\n{text}")
    if ext in VIDEO_EXT:
        # Whisper-style APIs accept mp4/webm/mp3/wav directly; keyframe captions can be added later.
        try:
            text = provider.transcribe(data, filename)
        except LLMError as e:
            raise UnsupportedSource(f"Could not transcribe the video/audio: {e}") from e
        kind = "audio" if ext in {".mp3", ".wav", ".m4a"} else "video"
        return ExtractedSource(kind, filename, f"[Transcript of {kind} narration]\n{text}")
    raise UnsupportedSource(f"Unsupported file type: {ext or filename}")


def _document(kind: str, filename: str, data: bytes) -> ExtractedSource:
    """PDF or DOCX text. Damaged, encrypted or text-less files become UnsupportedSource
    (a clear 415) instead of an unhandled 500."""
    try:
        text = _pdf(data) if kind == "pdf" else _docx(data)
    except Exception as e:  # pypdf and python-docx raise many unrelated types on bad input
        log.info("could not read %s %r: %s: %s", kind, filename, type(e).__name__, e)
        raise UnsupportedSource(_UNREADABLE[kind]) from e
    if not text.strip():
        raise UnsupportedSource(_NO_TEXT[kind])
    return ExtractedSource(kind, filename, text)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)
