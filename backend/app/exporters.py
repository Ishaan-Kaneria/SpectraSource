"""Turn an artefact's structured JSON into files operators can download and use."""

import html
import io
from collections.abc import Callable
from typing import Any

from app.schemas import OutputType

MEDIA = {
    "md": "text/markdown; charset=utf-8",
    "json": "application/json",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "srt": "application/x-subrip; charset=utf-8",
    "html": "text/html; charset=utf-8",
}


def _humanize(key: str) -> str:
    return key.replace("_", " ").capitalize()


# ---------- Markdown (every type) ----------


def to_markdown(content: dict[str, Any]) -> str:
    title = content.get("title") or content.get("hook") or "SpectraSource output"
    lines = [f"# {title}", ""]
    for k, v in content.items():
        if k == "title":
            continue
        lines += [f"## {_humanize(k)}", "", *_md_value(v), ""]
    return "\n".join(lines).rstrip() + "\n"


def _md_value(v: Any, indent: str = "") -> list[str]:
    if isinstance(v, list):
        out: list[str] = []
        for item in v:
            if isinstance(item, dict):
                parts = [f"**{_humanize(k)}:** {x}" for k, x in item.items() if not isinstance(x, (list, dict))]
                out.append(f"{indent}- " + "  \n  ".join(parts))
                for k, x in item.items():
                    if isinstance(x, (list, dict)):
                        out.append(f"{indent}  - **{_humanize(k)}:**")
                        out += _md_value(x, indent + "    ")
            else:
                out.append(f"{indent}- {item}")
        return out
    if isinstance(v, dict):
        return [f"{indent}- **{_humanize(k)}:** {x}" for k, x in v.items()]
    return [f"{indent}{v}"]


# ---------- DOCX (text-heavy types) ----------


def to_docx(content: dict[str, Any]) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(str(content.get("title") or content.get("hook") or "SpectraSource output"), 0)
    for k, v in content.items():
        if k == "title":
            continue
        if k == "sections" and isinstance(v, list):
            for sec in v:
                d.add_heading(str(sec.get("heading", "")), 1)
                for sk, sv in sec.items():
                    if sk != "heading":
                        d.add_paragraph(str(sv))
            continue
        d.add_heading(_humanize(k), 1)
        if isinstance(v, list):
            for item in v:
                if isinstance(item, dict):
                    d.add_paragraph(" | ".join(f"{_humanize(a)}: {b}" for a, b in item.items()), style="List Bullet")
                else:
                    d.add_paragraph(str(item), style="List Bullet")
        else:
            d.add_paragraph(str(v))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------- PPTX (presentation) ----------


def to_pptx(content: dict[str, Any]) -> bytes:
    from pptx import Presentation as Deck
    from pptx.util import Pt

    deck = Deck()
    title_slide = deck.slides.add_slide(deck.slide_layouts[0])
    title_slide.shapes.title.text = str(content.get("title", ""))
    title_slide.placeholders[1].text = str(content.get("subtitle", ""))

    for s in content.get("slides", []):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = str(s.get("title", ""))
        body = slide.placeholders[1].text_frame
        bullets = s.get("bullets", []) or [""]
        body.text = str(bullets[0])
        for b in bullets[1:]:
            body.add_paragraph().text = str(b)
        for p in body.paragraphs:
            for r in p.runs:
                r.font.size = Pt(20)
        notes = str(s.get("speaker_notes", ""))
        if s.get("visual_suggestion"):
            notes += f"\n\nVisual: {s['visual_suggestion']}"
        slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


# ---------- SRT (video subtitles) ----------


def to_srt(content: dict[str, Any]) -> str:
    blocks = []
    for i, line in enumerate(content.get("subtitles", []), 1):
        blocks.append(f"{i}\n{line.get('start', '')} --> {line.get('end', '')}\n{line.get('text', '')}\n")
    return "\n".join(blocks)


# ---------- HTML (infographic preview) ----------


def to_infographic_html(content: dict[str, Any]) -> str:
    e = lambda v: html.escape(str(v))  # noqa: E731
    palette = [c for c in content.get("color_palette", []) if isinstance(c, str) and c.startswith("#")]
    primary = palette[0] if palette else "#1e3a8a"
    accent = palette[1] if len(palette) > 1 else "#f59e0b"
    cards = "".join(
        f'<div class="card"><div class="stat">{e(s.get("key_stat_or_fact", ""))}</div>'
        f'<h3>{e(s.get("heading", ""))}</h3><p>{e(s.get("supporting_text", ""))}</p>'
        f'<small>Icon: {e(s.get("icon_suggestion", ""))}</small></div>'
        for s in content.get("sections", [])
    )
    return f"""<!doctype html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{e(content.get("title", ""))}</title>
<style>
body{{margin:0;font-family:system-ui,sans-serif;background:#f8fafc;color:#0f172a}}
.wrap{{max-width:900px;margin:0 auto;padding:24px}}
header{{background:{primary};color:#fff;padding:32px;border-radius:16px}}
header h1{{margin:0 0 8px;font-size:32px}} .key{{margin-top:16px;padding:12px 16px;background:{accent};color:#111;border-radius:10px;font-weight:600}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:16px;margin:24px 0}}
.card{{background:#fff;border-radius:14px;padding:20px;box-shadow:0 1px 3px rgba(0,0,0,.08)}}
.stat{{font-size:24px;font-weight:700;color:{primary}}} .card h3{{margin:8px 0 4px}} small{{color:#64748b}}
footer{{background:#0f172a;color:#fff;padding:16px 20px;border-radius:12px;font-weight:600}}
</style></head><body><div class="wrap">
<header><h1>{e(content.get("title", ""))}</h1><div>{e(content.get("subtitle", ""))}</div>
<div class="key">{e(content.get("key_message", ""))}</div></header>
<div class="grid">{cards}</div>
<footer>{e(content.get("call_to_action", ""))}</footer>
<p><small>Layout: {e(content.get("layout_recommendation", ""))}</small></p>
</div></body></html>"""


_TEXT_DOC = {"md", "docx", "json"}
EXPORTS: dict[OutputType, set[str]] = {
    OutputType.video: {"md", "srt", "docx", "json"},
    OutputType.linkedin: {"md", "json"},
    OutputType.twitter: {"md", "json"},
    OutputType.advisory: _TEXT_DOC,
    OutputType.infographic: {"html", "md", "json"},
    OutputType.executive_summary: _TEXT_DOC,
    OutputType.presentation: {"pptx", "md", "json"},
}

_RENDER: dict[str, Callable[[dict[str, Any]], str | bytes]] = {
    "md": to_markdown,
    "docx": to_docx,
    "pptx": to_pptx,
    "srt": to_srt,
    "html": to_infographic_html,
}


def render(output_type: OutputType, fmt: str, content: dict[str, Any]) -> bytes:
    import json

    if fmt not in EXPORTS[output_type]:
        raise ValueError(f"{output_type.value} cannot be exported as {fmt}")
    if fmt == "json":
        return json.dumps(content, indent=2, ensure_ascii=False).encode()
    out = _RENDER[fmt](content)
    return out.encode() if isinstance(out, str) else out
