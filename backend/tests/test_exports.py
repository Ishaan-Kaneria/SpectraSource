import io
import json

import pytest
from pptx import Presentation

from app.exporters import EXPORTS
from app.schemas import OutputType


@pytest.fixture(scope="module")
def artifacts(client):
    r = client.post(
        "/api/jobs",
        data={"output_types": json.dumps([t.value for t in OutputType]), "text": "Source text."},
    )
    job = client.get(f"/api/jobs/{r.json()['id']}").json()
    return {a["output_type"]: a for a in job["artifacts"]}


def test_every_advertised_format_downloads(client, artifacts):
    for t, a in artifacts.items():
        assert set(a["export_formats"]) == EXPORTS[OutputType(t)]
        for fmt in a["export_formats"]:
            r = client.get(f"/api/artifacts/{a['id']}/download", params={"fmt": fmt})
            assert r.status_code == 200, (t, fmt, r.text)
            assert r.content
            assert f'{t}-' in r.headers["content-disposition"]


def test_pptx_has_slides_and_notes(client, artifacts):
    r = client.get(f"/api/artifacts/{artifacts['presentation']['id']}/download", params={"fmt": "pptx"})
    deck = Presentation(io.BytesIO(r.content))
    assert len(deck.slides) >= 3  # title + content slides
    assert deck.slides[1].notes_slide.notes_text_frame.text


def test_srt_is_numbered(client, artifacts):
    r = client.get(f"/api/artifacts/{artifacts['video']['id']}/download", params={"fmt": "srt"})
    assert r.text.startswith("1\n") and "-->" in r.text


def test_unsupported_format_rejected(client, artifacts):
    r = client.get(f"/api/artifacts/{artifacts['linkedin']['id']}/download", params={"fmt": "pptx"})
    assert r.status_code == 400
